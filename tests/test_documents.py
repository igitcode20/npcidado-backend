from __future__ import annotations

import io

from docx import Document as WordDocument
from openpyxl import load_workbook

from app.services.document_builder import build_document
from app.services.ingestion import extract_content, strip_pdf_boilerplate
from app.services.text import chunk_text, index_terms


SOURCES = [
    {
        "document_id": "a" * 24,
        "title": "Manual de bienvenida",
        "filename": "manual.docx",
        "chunk_index": 0,
        "content": (
            "La empresa fue fundada en 2020. Su misión es organizar el conocimiento interno "
            "de las organizaciones.\n\nLos nuevos colaboradores deben completar el formulario "
            "de ingreso durante los primeros tres días."
        ),
    },
    {
        "document_id": "b" * 24,
        "title": "Políticas internas",
        "filename": "politicas.pdf",
        "chunk_index": 0,
        "content": (
            "Las políticas de seguridad exigen contraseñas de al menos doce caracteres. "
            "Los incidentes deben reportarse en las primeras 24 horas."
        ),
    },
]


def _build(fmt: str):
    return build_document(
        topic="Inducción de nuevos colaboradores",
        format=fmt,
        instructions="Enfocar en pasos y políticas",
        user_name="Ismael",
        sources=SOURCES,
    )


def test_build_docx() -> None:
    built = _build("docx")
    assert built.data[:2] == b"PK"
    document = WordDocument(io.BytesIO(built.data))
    body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    cover_text = "\n".join(
        paragraph.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    text = f"{cover_text}\n{body_text}"
    assert "Inducción de nuevos colaboradores" in text
    assert "Resumen ejecutivo" in text
    assert "Fuentes consultadas" in text
    assert built.filename.endswith(".docx")


def test_build_pdf() -> None:
    built = _build("pdf")
    assert built.data.startswith(b"%PDF")
    assert len(built.data) > 2000
    assert built.filename.endswith(".pdf")


def test_build_xlsx() -> None:
    built = _build("xlsx")
    workbook = load_workbook(io.BytesIO(built.data))
    assert {"Portada", "Ideas clave", "Contenido", "Fuentes"}.issubset(workbook.sheetnames)
    assert workbook["Contenido"].max_row >= 2
    assert built.filename.endswith(".xlsx")


def test_build_txt() -> None:
    built = _build("txt")
    text = built.data.decode("utf-8-sig")
    assert "NPICADO IA" in text
    assert "RESUMEN EJECUTIVO" in text
    assert "0 tokens de IA" in text


def test_chunking_and_terms() -> None:
    text = "Texto de prueba. " * 400
    chunks = chunk_text(text, size=300, overlap=50)
    assert len(chunks) > 5
    assert all(chunk.strip() for chunk in chunks)
    terms = index_terms("Políticas de seguridad y contraseñas")
    assert "politica" in terms
    assert "seguridad" in terms
    assert "contrasena" in terms


def test_strip_repeated_pdf_boilerplate() -> None:
    pages = [
        (1, "NPICADO IA\nPágina 1\nContenido real de la primera página.\nDocumento interno"),
        (2, "NPICADO IA\nPágina 2\nContenido real de la segunda página.\nDocumento interno"),
        (3, "NPICADO IA\nPágina 3\nContenido real de la tercera página.\nDocumento interno"),
    ]
    cleaned = strip_pdf_boilerplate(pages)
    text = "\n".join(content for _, content in cleaned)
    assert "Contenido real de la primera página" in text
    assert "Contenido real de la tercera página" in text
    assert "NPICADO IA" not in text
    assert "Página" not in text


def test_extract_plain_text_and_docx() -> None:
    plain = extract_content(".txt", "Contenido de prueba con acentos: ñandú.".encode("utf-8"))
    assert "ñandú" in plain.text

    source = WordDocument()
    source.add_paragraph("Un párrafo de prueba con información suficiente.")
    buffer = io.BytesIO()
    source.save(buffer)
    extracted = extract_content(".docx", buffer.getvalue())
    assert "párrafo de prueba" in extracted.text
