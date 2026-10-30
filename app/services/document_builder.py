from __future__ import annotations

import io
import re
import textwrap
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any
from xml.sax.saxutils import escape as xml_escape
from zoneinfo import ZoneInfo

from docx import Document as WordDocument
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    HRFlowable,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .text import (
    clean_display_text,
    content_tokens,
    excerpt,
    merge_overlapping_fragments,
    normalize_search_text,
    query_terms,
    slugify,
    split_sentences,
)

_FORMAT_META = {
    "docx": {
        "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "label": "Word",
    },
    "pdf": {"mime": "application/pdf", "label": "PDF"},
    "xlsx": {
        "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "label": "Excel",
    },
    "txt": {"mime": "text/plain; charset=utf-8", "label": "Texto"},
}
_MONTHS_ES = (
    "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
)


class EmptyKnowledgeError(Exception):
    pass


@dataclass
class DocumentContent:
    title: str
    topic: str
    instructions: str
    user_name: str
    date_text: str
    date_iso: str
    summary: str
    ideas: list[dict[str, str]] = field(default_factory=list)
    sections: list[dict[str, Any]] = field(default_factory=list)
    sources: list[dict[str, Any]] = field(default_factory=list)
    methodology_note: str = ""


@dataclass
class BuiltDocument:
    data: bytes
    filename: str
    mime_type: str
    format: str


def _format_date_es(value: datetime) -> str:
    return f"{value.day} de {_MONTHS_ES[value.month - 1]} de {value.year}"


def _clean_title(topic: str) -> str:
    title = re.sub(r"\s+", " ", clean_display_text(topic)).strip(" .,:;-")
    if not title:
        title = "Documento de conocimiento"
    if title.islower():
        title = title[0].upper() + title[1:]
    return title[:170]


def _split_blocks(text: str, max_block: int = 850) -> list[str]:
    blocks: list[str] = []
    for raw_block in re.split(r"\n\s*\n", text):
        raw_block = re.sub(r"[ \t]+", " ", raw_block).strip()
        raw_block = re.sub(
            r"^\[(?:Página|Pagina|Hoja|Sheet)[^\]]*\]\s*",
            "",
            raw_block,
            flags=re.IGNORECASE,
        ).strip()
        if not raw_block:
            continue
        if len(raw_block) <= max_block:
            blocks.append(raw_block)
            continue
        segments = re.split(r"(?<=[.!?])\s+|\n+", raw_block)
        buffer = ""
        for segment in segments:
            segment = segment.strip()
            if not segment:
                continue
            if len(buffer) + len(segment) + 1 <= max_block:
                buffer = f"{buffer} {segment}".strip()
            else:
                if buffer:
                    blocks.append(buffer)
                while len(segment) > max_block:
                    cut = segment.rfind(" ", 0, max_block)
                    cut = cut if cut > max_block * 0.6 else max_block
                    blocks.append(segment[:cut].strip())
                    segment = segment[cut:].strip()
                buffer = segment
        if buffer:
            blocks.append(buffer)
    return blocks


def _is_noise_block(block: str) -> bool:
    stripped = block.strip()
    if len(stripped) < 18:
        return True
    if re.fullmatch(r"\[(Página|Hoja|Sheet)[^\]]*\]", stripped, flags=re.IGNORECASE):
        return True
    return False


def _dedupe_blocks(blocks: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for block in blocks:
        key = re.sub(r"\W+", "", normalize_search_text(block))[:180]
        if not key or key in seen:
            continue
        seen.add(key)
        result.append(block)
    return result


def _build_groups(sources: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, dict[str, Any]] = {}
    for source in sources:
        doc_id = str(source.get("document_id", ""))
        if not doc_id:
            continue
        group = groups.setdefault(
            doc_id,
            {
                "document_id": doc_id,
                "title": source.get("title") or "Documento",
                "filename": source.get("filename", ""),
                "chunks": [],
                "source_numbers": [],
            },
        )
        if source.get("source_number"):
            group["source_numbers"].append(source["source_number"])
        group["chunks"].append(
            {
                "chunk_index": int(source.get("chunk_index", 0)),
                "content": source.get("content", ""),
            }
        )

    result: list[dict[str, Any]] = []
    for group in groups.values():
        group["chunks"].sort(key=lambda item: item["chunk_index"])
        merged = merge_overlapping_fragments([item["content"] for item in group["chunks"]])
        blocks = _dedupe_blocks(
            [block for block in _split_blocks(merged) if not _is_noise_block(block)]
        )
        if not blocks:
            continue
        selected: list[str] = []
        used_chars = 0
        for block in blocks:
            if used_chars >= 3400:
                break
            if len(block) > 950:
                block = excerpt(block, 900)
            selected.append(block)
            used_chars += len(block)
        if selected:
            group["paragraphs"] = selected
            group.pop("chunks", None)
            result.append(group)
    return result


def _extract_ideas(groups: list[dict[str, Any]], query: str, limit: int = 7) -> list[dict[str, str]]:
    terms = set(query_terms(query))
    scored: list[tuple[float, str, str]] = []
    seen: set[str] = set()
    for group in groups:
        for paragraph in group["paragraphs"]:
            for sentence in split_sentences(paragraph):
                clean = sentence.strip()
                if len(clean) < 45 or len(clean) > 320:
                    continue
                key = re.sub(r"\W+", "", normalize_search_text(clean))[:140]
                if key in seen:
                    continue
                seen.add(key)
                sentence_terms = set(content_tokens(clean))
                overlap = len(terms & sentence_terms)
                score = overlap * 2.0 + min(len(clean), 220) / 300
                scored.append((score, clean, group["title"]))
    scored.sort(key=lambda item: item[0], reverse=True)
    ideas: list[dict[str, str]] = []
    for _, sentence, source_title in scored[:limit]:
        ideas.append({"text": sentence, "source": source_title})
    return ideas


def prepare_content(
    *,
    topic: str,
    instructions: str,
    user_name: str,
    sources: list[dict[str, Any]],
    timezone: str = "America/Managua",
    now: datetime | None = None,
) -> DocumentContent:
    groups = _build_groups(sources)
    if not groups:
        raise EmptyKnowledgeError(
            "No encontré contenido relacionado en la biblioteca. Sube documentos sobre este tema o prueba con otras palabras."
        )

    local_now = (now or datetime.now(ZoneInfo(timezone))).astimezone(ZoneInfo(timezone))
    title = _clean_title(topic)
    instructions = clean_display_text(instructions)
    user_name = clean_display_text(user_name)[:80] or "Usuario de NPICADO IA"
    query = f"{topic} {instructions}".strip()
    ideas = _extract_ideas(groups, query)
    summary_parts = [idea["text"] for idea in ideas[:4]]
    if summary_parts:
        summary = " ".join(summary_parts)
    else:
        summary = excerpt(groups[0]["paragraphs"][0], 700)
    summary = clean_display_text(summary)

    sections: list[dict[str, Any]] = []
    for index, group in enumerate(groups, start=1):
        sections.append(
            {
                "number": index,
                "title": group["title"],
                "filename": group["filename"],
                "paragraphs": group["paragraphs"],
            }
        )

    source_list = []
    for index, group in enumerate(groups, start=1):
        source_list.append(
            {
                "number": index,
                "title": group["title"],
                "filename": group["filename"],
            }
        )

    methodology_note = (
        "Este documento fue compilado automáticamente por Python con fragmentos textuales "
        "de la biblioteca privada de NPICADO IA. La IA no redactó este archivo y no se "
        "consumieron tokens de Groq para generarlo. La información conserva el contenido "
        "de las fuentes citadas."
    )
    return DocumentContent(
        title=title,
        topic=topic.strip(),
        instructions=instructions.strip(),
        user_name=user_name,
        date_text=_format_date_es(local_now),
        date_iso=local_now.isoformat(),
        summary=summary,
        ideas=ideas,
        sections=sections,
        sources=source_list,
        methodology_note=methodology_note,
    )


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------


def _shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _set_cell_margins(cell, top=280, start=340, bottom=280, end=340) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(fld_end)


def _render_docx(content: DocumentContent) -> bytes:
    document = WordDocument()
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x1F, 0x2A, 0x3A)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    for style_name, color, size in (
        ("Heading 1", "0B6E69", 18),
        ("Heading 2", "123B5C", 13),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(16)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True

    document.core_properties.title = content.title
    document.core_properties.author = "NPICADO IA"
    document.core_properties.subject = "Documento de conocimiento compilado por Python"
    document.core_properties.keywords = "NPICADO IA, biblioteca, documento"
    document.core_properties.comments = "Generado sin consumo de tokens de IA."

    # Portada oscura en una tabla de una celda.
    cover = document.add_table(rows=1, cols=1)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.autofit = False
    cell = cover.cell(0, 0)
    cell.width = Cm(16.4)
    cover.rows[0].height = Cm(22.2)
    _shade_cell(cell, "0A1524")
    _set_cell_margins(cell, top=650, start=650, bottom=500, end=650)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    brand = cell.paragraphs[0]
    brand.paragraph_format.space_after = Pt(4)
    run = brand.add_run("NPICADO IA")
    run.font.name = "Consolas"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x5E, 0xEA, 0xD4)

    rule = cell.add_paragraph()
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(68)
    p_pr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "2DD4BF")
    borders.append(bottom)
    p_pr.append(borders)

    title_paragraph = cell.add_paragraph()
    title_paragraph.paragraph_format.space_after = Pt(14)
    title_run = title_paragraph.add_run(content.title)
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(29)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0xF4, 0xF8, 0xFC)

    subtitle = cell.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(80)
    subtitle_run = subtitle.add_run("Documento de conocimiento")
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = RGBColor(0x9F, 0xB4, 0xC9)

    meta = cell.add_paragraph()
    meta.paragraph_format.space_after = Pt(2)
    meta_run = meta.add_run(f"Preparado para  {content.user_name}")
    meta_run.font.size = Pt(10)
    meta_run.font.color.rgb = RGBColor(0xCB, 0xD8, 0xE5)

    date_p = cell.add_paragraph()
    date_p.paragraph_format.space_after = Pt(18)
    date_run = date_p.add_run(content.date_text)
    date_run.font.size = Pt(10)
    date_run.font.color.rgb = RGBColor(0xCB, 0xD8, 0xE5)

    stamp = cell.add_paragraph()
    stamp.paragraph_format.space_after = Pt(0)
    stamp_run = stamp.add_run("Compilado por Python  ·  0 tokens de IA")
    stamp_run.font.name = "Consolas"
    stamp_run.font.size = Pt(8.5)
    stamp_run.font.color.rgb = RGBColor(0x5E, 0xEA, 0xD4)

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Pie de página.
    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run("NPICADO IA  ·  Documento de conocimiento  ·  página ")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(0x6B, 0x7C, 0x90)
    _add_page_number(footer_paragraph)

    document.add_heading("Resumen ejecutivo", level=1)
    for paragraph_text in _split_blocks(content.summary, 1400):
        p = document.add_paragraph(paragraph_text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if content.instructions:
        document.add_heading("Enfoque solicitado", level=2)
        p = document.add_paragraph(content.instructions)
        p.paragraph_format.left_indent = Cm(0.4)
        for run in p.runs:
            run.italic = True
            run.font.color.rgb = RGBColor(0x47, 0x5A, 0x70)

    if content.ideas:
        document.add_heading("Ideas clave", level=1)
        for idea in content.ideas:
            p = document.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(5)
            p.add_run(idea["text"])
            source_run = p.add_run(f"  — {idea['source']}")
            source_run.italic = True
            source_run.font.size = Pt(8.5)
            source_run.font.color.rgb = RGBColor(0x5F, 0x74, 0x89)

    document.add_heading("Desarrollo", level=1)
    for section in content.sections:
        document.add_heading(f"{section['number']}. {section['title']}", level=2)
        for paragraph_text in section["paragraphs"]:
            p = document.add_paragraph(paragraph_text)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    document.add_heading("Fuentes consultadas", level=1)
    source_table = document.add_table(rows=1, cols=3)
    source_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    source_table.style = "Light Shading Accent 1"
    headers = ("#", "Documento", "Archivo")
    for index, header in enumerate(headers):
        cell = source_table.rows[0].cells[index]
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)

    for source in content.sources:
        cells = source_table.add_row().cells
        cells[0].text = str(source["number"])
        cells[1].text = source["title"]
        cells[2].text = source["filename"]
        for cell in cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)

    document.add_paragraph()
    note = document.add_paragraph(content.methodology_note)
    note.paragraph_format.left_indent = Cm(0.3)
    note.paragraph_format.space_before = Pt(10)
    for run in note.runs:
        run.italic = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x5B, 0x6B, 0x7C)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _pdf_safe(value: str) -> str:
    value = clean_display_text(value)
    value = value.replace("→", "->").replace("←", "<-").replace("•", "·")
    return value.encode("cp1252", errors="replace").decode("cp1252")


def _pdf_paragraph(value: str) -> str:
    return xml_escape(_pdf_safe(value)).replace("\n", "<br/>")


def _pdf_page(canvas, doc) -> None:
    canvas.saveState()
    page_number = canvas.getPageNumber()
    width, height = A4
    if page_number == 1:
        canvas.setFillColor(colors.HexColor("#07111F"))
        canvas.rect(0, 0, width, height, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor("#123047"))
        canvas.circle(width - 25, height - 40, 170, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor("#0B5A57"))
        canvas.circle(width + 10, -20, 150, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor("#5EEAD4"))
        canvas.rect(2 * cm, height - 2.5 * cm, 2.6 * cm, 0.09 * cm, fill=1, stroke=0)
    else:
        canvas.setStrokeColor(colors.HexColor("#DCE6EE"))
        canvas.setLineWidth(0.5)
        canvas.line(2 * cm, height - 1.5 * cm, width - 2 * cm, height - 1.5 * cm)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#6C8196"))
        canvas.drawString(2 * cm, height - 1.2 * cm, "NPICADO IA  ·  DOCUMENTO DE CONOCIMIENTO")
        canvas.setStrokeColor(colors.HexColor("#DCE6EE"))
        canvas.line(2 * cm, 1.45 * cm, width - 2 * cm, 1.45 * cm)
        canvas.drawString(2 * cm, 0.95 * cm, f"Generado el {doc._npicado_date}")
        canvas.drawRightString(width - 2 * cm, 0.95 * cm, f"Página {page_number}")
    canvas.restoreState()


def _render_pdf(content: DocumentContent) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=2.1 * cm,
        leftMargin=2.1 * cm,
        topMargin=2.0 * cm,
        bottomMargin=1.7 * cm,
        title=content.title,
        author="NPICADO IA",
        subject="Documento de conocimiento compilado por Python",
        creator="NPICADO IA",
    )
    doc._npicado_date = content.date_text  # type: ignore[attr-defined]

    base = getSampleStyleSheet()
    brand_style = ParagraphStyle(
        "NPBrand",
        parent=base["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#5EEAD4"),
        spaceAfter=8,
    )
    cover_title = ParagraphStyle(
        "NPCoverTitle",
        parent=base["Title"],
        fontName="Helvetica-Bold",
        fontSize=29,
        leading=35,
        textColor=colors.HexColor("#F3F8FC"),
        alignment=TA_LEFT,
        spaceAfter=18,
    )
    cover_sub = ParagraphStyle(
        "NPCoverSub",
        parent=base["Normal"],
        fontName="Helvetica",
        fontSize=13,
        leading=18,
        textColor=colors.HexColor("#A9BDD0"),
        spaceAfter=6,
    )
    cover_meta = ParagraphStyle(
        "NPCoverMeta",
        parent=base["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=15,
        textColor=colors.HexColor("#C7D5E2"),
    )
    cover_stamp = ParagraphStyle(
        "NPCoverStamp",
        parent=cover_meta,
        fontName="Courier-Bold",
        fontSize=8,
        textColor=colors.HexColor("#5EEAD4"),
    )
    h1 = ParagraphStyle(
        "NPH1",
        parent=base["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=17,
        leading=21,
        textColor=colors.HexColor("#0B5E59"),
        spaceBefore=16,
        spaceAfter=8,
        keepWithNext=1,
    )
    h2 = ParagraphStyle(
        "NPH2",
        parent=base["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#183D5A"),
        spaceBefore=13,
        spaceAfter=6,
        keepWithNext=1,
    )
    body = ParagraphStyle(
        "NPBody",
        parent=base["BodyText"],
        fontName="Helvetica",
        fontSize=9.8,
        leading=15,
        textColor=colors.HexColor("#26384A"),
        alignment=TA_JUSTIFY,
        spaceAfter=7,
    )
    small = ParagraphStyle(
        "NPSmall",
        parent=body,
        fontSize=8.2,
        leading=11.5,
        textColor=colors.HexColor("#667A8E"),
    )
    source_cell = ParagraphStyle(
        "NPSource",
        parent=base["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#26384A"),
    )
    source_head = ParagraphStyle(
        "NPSourceHead",
        parent=source_cell,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    story: list[Any] = []
    story.append(Spacer(1, 3.7 * cm))
    story.append(Paragraph("NPICADO IA", brand_style))
    story.append(HRFlowable(width="17%", thickness=2, color=colors.HexColor("#2DD4BF"), spaceAfter=28))
    story.append(Paragraph(_pdf_paragraph(content.title), cover_title))
    story.append(Paragraph("Documento de conocimiento", cover_sub))
    story.append(Spacer(1, 3.5 * cm))
    story.append(Paragraph(f"Preparado para <b>{_pdf_paragraph(content.user_name)}</b>", cover_meta))
    story.append(Paragraph(_pdf_paragraph(content.date_text), cover_meta))
    story.append(Paragraph(f"{len(content.sources)} fuente(s) de la biblioteca", cover_meta))
    story.append(Spacer(1, 2.0 * cm))
    story.append(Paragraph("COMPILADO POR PYTHON  ·  0 TOKENS DE IA", cover_stamp))
    story.append(PageBreak())

    story.append(Paragraph("Resumen ejecutivo", h1))
    for block in _split_blocks(content.summary, 1400):
        story.append(Paragraph(_pdf_paragraph(block), body))

    if content.instructions:
        story.append(Paragraph("Enfoque solicitado", h2))
        story.append(Paragraph(f"<i>{_pdf_paragraph(content.instructions)}</i>", body))

    if content.ideas:
        story.append(Paragraph("Ideas clave", h1))
        items = [
            ListItem(
                Paragraph(
                    f"{_pdf_paragraph(idea['text'])} <font color='#6B8196' size='7.5'>— {_pdf_paragraph(idea['source'])}</font>",
                    body,
                ),
                leftIndent=12,
            )
            for idea in content.ideas
        ]
        story.append(
            ListFlowable(
                items,
                bulletType="bullet",
                start="•",
                leftIndent=14,
                bulletFontName="Helvetica",
                bulletFontSize=8,
                bulletColor=colors.HexColor("#0F766E"),
                spaceAfter=6,
            )
        )

    story.append(Paragraph("Desarrollo", h1))
    for section in content.sections:
        story.append(Paragraph(f"{section['number']}. {_pdf_paragraph(section['title'])}", h2))
        for block in section["paragraphs"]:
            story.append(Paragraph(_pdf_paragraph(block), body))

    story.append(Paragraph("Fuentes consultadas", h1))
    table_data: list[list[Any]] = [
        [
            Paragraph("#", source_head),
            Paragraph("Documento", source_head),
            Paragraph("Archivo", source_head),
        ]
    ]
    for source in content.sources:
        table_data.append(
            [
                Paragraph(str(source["number"]), source_cell),
                Paragraph(_pdf_paragraph(source["title"]), source_cell),
                Paragraph(_pdf_paragraph(source["filename"]), source_cell),
            ]
        )
    sources_table = Table(table_data, colWidths=[1.0 * cm, 8.7 * cm, 6.7 * cm], repeatRows=1)
    sources_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A5A56")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#DCE7EE")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F9FB")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(sources_table)
    story.append(Spacer(1, 0.35 * cm))
    story.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#DDE7EE")))
    story.append(Spacer(1, 0.15 * cm))
    story.append(Paragraph(_pdf_paragraph(content.methodology_note), small))

    doc.build(story, onFirstPage=_pdf_page, onLaterPages=_pdf_page)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def _excel_value(value: str) -> str:
    clean = clean_display_text(value)
    if clean.startswith(("=", "+", "-", "@", "\t", "\r")):
        clean = "'" + clean
    return clean


def _style_header_row(sheet, row: int, columns: int) -> None:
    fill = PatternFill("solid", fgColor="0A5A56")
    font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    for column in range(1, columns + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 24


def _render_xlsx(content: DocumentContent) -> bytes:
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Portada"
    cover.sheet_view.showGridLines = False

    dark = PatternFill("solid", fgColor="07111F")
    teal = PatternFill("solid", fgColor="2DD4BF")
    for row in cover.iter_rows(min_row=1, max_row=36, min_col=1, max_col=8):
        for cell in row:
            cell.fill = dark
    for column in range(1, 9):
        cover.column_dimensions[get_column_letter(column)].width = 14
    for row in range(1, 37):
        cover.row_dimensions[row].height = 18

    cover.merge_cells("B3:G4")
    brand = cover["B3"]
    brand.value = "NPICADO IA"
    brand.font = Font(name="Consolas", size=12, bold=True, color="5EEAD4")
    brand.alignment = Alignment(vertical="center")

    cover.merge_cells("B6:G7")
    title = cover["B6"]
    title.value = _excel_value(content.title)
    title.font = Font(name="Calibri", size=25, bold=True, color="F3F8FC")
    title.alignment = Alignment(wrap_text=True, vertical="center")

    for column in range(2, 8):
        cover.cell(row=8, column=column).fill = teal
    cover.row_dimensions[8].height = 3

    cover.merge_cells("B10:G10")
    cover["B10"].value = "DOCUMENTO DE CONOCIMIENTO"
    cover["B10"].font = Font(name="Calibri", size=11, color="9FB4C9")

    cover.merge_cells("B13:G13")
    cover["B13"].value = f"Preparado para: {content.user_name}"
    cover["B13"].font = Font(name="Calibri", size=10, color="E1EAF3")
    cover.merge_cells("B14:G14")
    cover["B14"].value = content.date_text
    cover["B14"].font = Font(name="Calibri", size=10, color="E1EAF3")
    cover.merge_cells("B15:G15")
    cover["B15"].value = f"{len(content.sources)} fuente(s) de la biblioteca"
    cover["B15"].font = Font(name="Calibri", size=10, color="E1EAF3")

    cover.merge_cells("B19:G21")
    cover["B19"].value = "COMPILADO POR PYTHON  ·  0 TOKENS DE IA"
    cover["B19"].font = Font(name="Consolas", size=9, bold=True, color="5EEAD4")
    cover["B19"].alignment = Alignment(vertical="center")

    ideas_sheet = workbook.create_sheet("Ideas clave")
    ideas_sheet.sheet_view.showGridLines = False
    headers = ["#", "Idea clave", "Fuente"]
    ideas_sheet.append(headers)
    _style_header_row(ideas_sheet, 1, 3)
    for index, idea in enumerate(content.ideas, start=1):
        ideas_sheet.append([index, _excel_value(idea["text"]), _excel_value(idea["source"])])
        row = ideas_sheet.max_row
        ideas_sheet.row_dimensions[row].height = 54
        for column in range(1, 4):
            ideas_sheet.cell(row=row, column=column).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
    ideas_sheet.column_dimensions["A"].width = 5
    ideas_sheet.column_dimensions["B"].width = 90
    ideas_sheet.column_dimensions["C"].width = 34
    ideas_sheet.freeze_panes = "A2"
    if ideas_sheet.max_row > 1:
        ideas_sheet.auto_filter.ref = f"A1:C{ideas_sheet.max_row}"

    content_sheet = workbook.create_sheet("Contenido")
    content_sheet.sheet_view.showGridLines = False
    content_sheet.append(["#", "Documento", "Archivo", "Texto"])
    _style_header_row(content_sheet, 1, 4)
    row_number = 0
    for section in content.sections:
        for paragraph in section["paragraphs"]:
            row_number += 1
            content_sheet.append(
                [
                    row_number,
                    _excel_value(section["title"]),
                    _excel_value(section["filename"]),
                    _excel_value(paragraph),
                ]
            )
            row = content_sheet.max_row
            content_sheet.row_dimensions[row].height = max(46, min(120, len(paragraph) / 2))
            for column in range(1, 5):
                content_sheet.cell(row=row, column=column).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )
    content_sheet.column_dimensions["A"].width = 5
    content_sheet.column_dimensions["B"].width = 40
    content_sheet.column_dimensions["C"].width = 38
    content_sheet.column_dimensions["D"].width = 100
    content_sheet.freeze_panes = "A2"
    if content_sheet.max_row > 1:
        content_sheet.auto_filter.ref = f"A1:D{content_sheet.max_row}"

    sources_sheet = workbook.create_sheet("Fuentes")
    sources_sheet.sheet_view.showGridLines = False
    sources_sheet.append(["#", "Documento", "Archivo"])
    _style_header_row(sources_sheet, 1, 3)
    for source in content.sources:
        sources_sheet.append(
            [source["number"], _excel_value(source["title"]), _excel_value(source["filename"])]
        )
        sources_sheet.row_dimensions[sources_sheet.max_row].height = 30
        for column in range(1, 4):
            sources_sheet.cell(row=sources_sheet.max_row, column=column).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
    sources_sheet.column_dimensions["A"].width = 5
    sources_sheet.column_dimensions["B"].width = 55
    sources_sheet.column_dimensions["C"].width = 55
    sources_sheet.freeze_panes = "A2"

    notes_row = sources_sheet.max_row + 3
    sources_sheet.merge_cells(start_row=notes_row, start_column=1, end_row=notes_row, end_column=3)
    notes_cell = sources_sheet.cell(row=notes_row, column=1)
    notes_cell.value = _excel_value(content.methodology_note)
    notes_cell.font = Font(name="Calibri", size=9, italic=True, color="5B6B7C")
    notes_cell.alignment = Alignment(wrap_text=True, vertical="top")
    sources_sheet.row_dimensions[notes_row].height = 58

    for sheet in (ideas_sheet, content_sheet, sources_sheet):
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    workbook.properties.creator = "NPICADO IA"
    workbook.properties.title = content.title
    workbook.properties.subject = "Documento de conocimiento compilado por Python"
    workbook.properties.description = "Generado sin consumo de tokens de IA."
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Texto plano
# ---------------------------------------------------------------------------


def _render_txt(content: DocumentContent) -> bytes:
    line = "=" * 78
    thin = "-" * 78
    lines: list[str] = [
        "NPICADO IA",
        "DOCUMENTO DE CONOCIMIENTO",
        line,
        content.title,
        line,
        f"Preparado para : {content.user_name}",
        f"Fecha           : {content.date_text}",
        f"Fuentes         : {len(content.sources)}",
        "Generado por    : Python (0 tokens de IA)",
        "",
        "RESUMEN EJECUTIVO",
        thin,
        "",
    ]
    for paragraph in _split_blocks(content.summary, 1000):
        lines.extend(textwrap.wrap(paragraph, width=96) or [""])
        lines.append("")

    if content.instructions:
        lines.extend(["ENFOQUE SOLICITADO", thin, ""])
        lines.extend(textwrap.wrap(content.instructions, width=96))
        lines.append("")

    if content.ideas:
        lines.extend(["IDEAS CLAVE", thin, ""])
        for index, idea in enumerate(content.ideas, start=1):
            wrapped = textwrap.wrap(idea["text"], width=94) or [""]
            lines.append(f"{index}. {wrapped[0]}")
            lines.extend(f"   {part}" for part in wrapped[1:])
            lines.append(f"   Fuente: {idea['source']}")
            lines.append("")
        lines.append("")

    lines.extend(["DESARROLLO", thin, ""])
    for section in content.sections:
        lines.extend([f"{section['number']}. {section['title'].upper()}", thin, ""])
        for paragraph in section["paragraphs"]:
            lines.extend(textwrap.wrap(paragraph, width=96) or [""])
            lines.append("")

    lines.extend(["FUENTES CONSULTADAS", thin, ""])
    for source in content.sources:
        lines.append(f"[{source['number']}] {source['title']}")
        lines.append(f"    Archivo: {source['filename']}")
    lines.extend(["", thin, "NOTA METODOLÓGICA", ""])
    lines.extend(textwrap.wrap(content.methodology_note, width=96))
    lines.extend(["", line, "NPICADO IA · Biblioteca privada · Documento compilado por Python", line, ""])
    return "\n".join(lines).encode("utf-8-sig")


def build_document(
    *,
    topic: str,
    format: str,
    instructions: str,
    user_name: str,
    sources: list[dict[str, Any]],
    timezone: str = "America/Managua",
) -> BuiltDocument:
    if format not in _FORMAT_META:
        raise ValueError("Formato no soportado.")
    content = prepare_content(
        topic=topic,
        instructions=instructions,
        user_name=user_name,
        sources=sources,
        timezone=timezone,
    )
    renderers = {
        "docx": _render_docx,
        "pdf": _render_pdf,
        "xlsx": _render_xlsx,
        "txt": _render_txt,
    }
    data = renderers[format](content)
    local_now = datetime.now(ZoneInfo(timezone))
    filename = (
        f"NPICADO_{slugify(content.title, 52)}_{local_now.strftime('%Y%m%d_%H%M%S')}."
        f"{format}"
    )
    return BuiltDocument(
        data=data,
        filename=filename,
        mime_type=_FORMAT_META[format]["mime"],
        format=format,
    )
