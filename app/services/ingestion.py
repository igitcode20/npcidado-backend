from __future__ import annotations

import csv
import hashlib
import io
import logging
import math
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader
from pymongo.errors import DuplicateKeyError

from ..bootstrap import utcnow
from ..config import Settings
from ..database import Database
from ..serializers import public_document
from ..services.text import (
    chunk_text,
    clean_display_text,
    excerpt,
    index_terms,
    safe_filename,
)

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".pdf", ".docx", ".xlsx"}
EXTENSION_MIME = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class IngestionError(Exception):
    def __init__(self, message: str, status_code: int = 422) -> None:
        super().__init__(message)
        self.message = message
        self.status_code = status_code


@dataclass
class ExtractedContent:
    text: str
    pages: int | None = None


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _boilerplate_key(line: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\d+", "#", line.lower())).strip()


def strip_pdf_boilerplate(page_texts: list[tuple[int, str]]) -> list[tuple[int, str]]:
    """Elimina encabezados/pies repetidos (empresa, «página N», fecha…).

    Solo considera líneas de los extremos de cada página que se repiten en al
    menos la mitad de las páginas. Así se reduce ruido sin borrar contenido real.
    """
    if len(page_texts) < 2:
        return page_texts
    pages = [(number, [line.strip() for line in text.splitlines() if line.strip()]) for number, text in page_texts]
    counts: Counter[str] = Counter()
    for _, lines in pages:
        edge_lines = {_boilerplate_key(line) for line in lines[:3] + lines[-3:]}
        for key in edge_lines:
            if 3 <= len(key) <= 140 and not key.isdigit():
                counts[key] += 1
    threshold = max(2, math.ceil(len(pages) / 2))
    repeated = {key for key, count in counts.items() if count >= threshold}
    if not repeated:
        return page_texts

    cleaned: list[tuple[int, str]] = []
    for number, lines in pages:
        kept = [line for line in lines if _boilerplate_key(line) not in repeated]
        text = re.sub(r"\n{3,}", "\n\n", "\n".join(kept)).strip()
        if text:
            cleaned.append((number, text))
    return cleaned


def _extract_pdf(data: bytes) -> ExtractedContent:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise IngestionError("El PDF está protegido con contraseña y no se puede leer.") from exc
        raw_pages: list[tuple[int, str]] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                raw_pages.append((page_number, text))
        if not raw_pages:
            raise IngestionError("El PDF no tiene texto seleccionable. Si es un escaneo, primero aplica OCR.")
        cleaned_pages = strip_pdf_boilerplate(raw_pages)
        pages = [f"[Página {number}]\n{text}" for number, text in cleaned_pages]
        return ExtractedContent("\n\n".join(pages), pages=len(reader.pages))
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("No se pudo leer el PDF. Verifica que el archivo no esté dañado.") from exc


def _extract_docx(data: bytes) -> ExtractedContent:
    try:
        document = DocxDocument(io.BytesIO(data))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    parts.append(" | ".join(values))
        return ExtractedContent("\n\n".join(parts))
    except Exception as exc:
        raise IngestionError("No se pudo leer el DOCX. Verifica que sea un documento de Word válido.") from exc


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().replace("\n", " ")


def _extract_xlsx(data: bytes) -> ExtractedContent:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts: list[str] = []
        total_chars = 0
        for sheet in workbook.worksheets:
            parts.append(f"[Hoja: {sheet.title}]")
            total_chars += len(sheet.title) + 10
            for row in sheet.iter_rows(values_only=True):
                values = [_cell_text(value) for value in row]
                values = [value for value in values if value]
                if values:
                    line = " | ".join(values)
                    parts.append(line)
                    total_chars += len(line) + 1
                if total_chars > 2_000_000:
                    raise IngestionError("El Excel contiene demasiado texto. Divide el archivo antes de subirlo.")
        workbook.close()
        return ExtractedContent("\n".join(parts))
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("No se pudo leer el Excel. Verifica que sea un archivo .xlsx válido.") from exc


def _extract_csv(data: bytes) -> ExtractedContent:
    text = _decode_text(data)
    try:
        sample = text[:4096]
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        rows = csv.reader(io.StringIO(text), dialect)
        lines: list[str] = []
        for index, row in enumerate(rows):
            if index > 20000:
                break
            values = [value.strip().replace("\n", " ") for value in row if value.strip()]
            if values:
                lines.append(" | ".join(values))
        return ExtractedContent("\n".join(lines))
    except csv.Error:
        return ExtractedContent(text)


def extract_content(extension: str, data: bytes) -> ExtractedContent:
    if extension in {".txt", ".md"}:
        return ExtractedContent(_decode_text(data))
    if extension == ".csv":
        return _extract_csv(data)
    if extension == ".pdf":
        return _extract_pdf(data)
    if extension == ".docx":
        return _extract_docx(data)
    if extension == ".xlsx":
        return _extract_xlsx(data)
    raise IngestionError("Formato no soportado. Usa PDF, DOCX, XLSX, TXT, MD o CSV.")


def _fallback_title(filename: str) -> str:
    stem = Path(filename).stem.replace("_", " ").replace("-", " ")
    stem = re.sub(r"\s+", " ", stem).strip()
    return (stem or "Documento sin título")[:120]


def _clean_tags(tags: list[str] | str | None) -> list[str]:
    if not tags:
        return []
    if isinstance(tags, str):
        raw = re.split(r"[,;]", tags)
    else:
        raw = tags
    result: list[str] = []
    for tag in raw:
        clean = re.sub(r"\s+", " ", str(tag)).strip()[:32]
        if clean and clean.lower() not in {item.lower() for item in result}:
            result.append(clean)
    return result[:10]


def ingest_document(
    database: Database,
    settings: Settings,
    *,
    filename: str,
    data: bytes,
    title: str | None,
    tags: list[str] | str | None,
    uploaded_by: dict[str, Any],
) -> dict[str, Any]:
    display_filename = safe_filename(filename)
    extension = Path(display_filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise IngestionError("Formato no soportado. Puedes subir PDF, DOCX, XLSX, TXT, MD o CSV.")
    if not data:
        raise IngestionError("El archivo está vacío.")
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise IngestionError(f"El archivo supera el máximo de {settings.max_upload_mb} MB.", 413)

    sha256 = hashlib.sha256(data).hexdigest()
    duplicate = database.documents.find_one({"sha256": sha256}, {"title": 1})
    if duplicate:
        raise IngestionError(f"Este archivo ya está en la biblioteca como «{duplicate.get('title', 'documento')}».", 409)

    extracted = extract_content(extension, data)
    text = clean_display_text(extracted.text)
    if len(text) < 40:
        raise IngestionError("El documento no contiene texto suficiente para alimentar la biblioteca.")
    if len(text) > settings.max_document_chars:
        raise IngestionError(
            f"El documento supera el límite de {settings.max_document_chars:,} caracteres. "
            "Divídelo en partes más pequeñas."
        )

    chunks = chunk_text(text)
    if not chunks:
        raise IngestionError("No se pudo extraer contenido útil del documento.")

    now = utcnow()
    clean_title = re.sub(r"\s+", " ", (title or "").strip())[:160] or _fallback_title(display_filename)
    clean_tags = _clean_tags(tags)
    document = {
        "title": clean_title,
        "filename": display_filename,
        "extension": extension.lstrip("."),
        "mime_type": EXTENSION_MIME.get(extension, "application/octet-stream"),
        "size_bytes": len(data),
        "char_count": len(text),
        "page_count": extracted.pages,
        "chunk_count": len(chunks),
        "tags": clean_tags,
        "preview": excerpt(text, 360),
        "sha256": sha256,
        "uploaded_by": uploaded_by["_id"],
        "uploaded_by_name": uploaded_by.get("full_name") or uploaded_by["username"],
        "uploaded_at": now,
        "updated_at": now,
        "status": "ready",
    }

    try:
        result = database.documents.insert_one(document)
    except DuplicateKeyError as exc:
        raise IngestionError("Este archivo ya está en la biblioteca.", 409) from exc

    document_id = result.inserted_id
    chunk_docs = []
    for index, content in enumerate(chunks):
        chunk_docs.append(
            {
                "document_id": document_id,
                "chunk_index": index,
                "content": content,
                "terms": index_terms(f"{clean_title} {content}"),
                "title": clean_title,
                "filename": display_filename,
                "extension": extension.lstrip("."),
                "created_at": now,
            }
        )
    try:
        database.chunks.insert_many(chunk_docs)
    except Exception:
        database.chunks.delete_many({"document_id": document_id})
        database.documents.delete_one({"_id": document_id})
        logger.exception("Fallo insertando fragmentos del documento %s", document_id)
        raise IngestionError("No se pudieron guardar los fragmentos del documento.", 500) from None

    document["_id"] = document_id
    logger.info(
        "Documento '%s' ingestado: %s fragmentos, %s caracteres.",
        clean_title,
        len(chunks),
        len(text),
    )
    return public_document(document)


def delete_document(database: Database, document_id) -> bool:
    deleted = database.documents.delete_one({"_id": document_id})
    database.chunks.delete_many({"document_id": document_id})
    return deleted.deleted_count > 0
